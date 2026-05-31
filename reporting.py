from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Sequence

try:
    from docx import Document  # type: ignore[import-not-found]
    from docx.shared import Inches  # type: ignore[import-not-found]
except Exception as exc:  # pragma: no cover - optional dependency guard
    Document = None  # type: ignore[assignment]
    Inches = None  # type: ignore[assignment]
    _DOCX_IMPORT_ERROR = exc

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
except Exception as exc:  # pragma: no cover - optional dependency guard
    colors = None  # type: ignore[assignment]
    letter = None  # type: ignore[assignment]
    getSampleStyleSheet = None  # type: ignore[assignment]
    ParagraphStyle = None  # type: ignore[assignment]
    inch = None  # type: ignore[assignment]
    Paragraph = None  # type: ignore[assignment]
    SimpleDocTemplate = None  # type: ignore[assignment]
    Spacer = None  # type: ignore[assignment]
    Table = None  # type: ignore[assignment]
    TableStyle = None  # type: ignore[assignment]
    _PDF_IMPORT_ERROR = exc

from essay_pipeline import EssayEvaluation


def build_docx_report(
    submission_id: str,
    student_name: str,
    prompt: str,
    essay_text: str,
    evaluation: EssayEvaluation,
    references: Sequence[str] | None = None,
) -> BytesIO:
    if Document is None:
        raise RuntimeError(
            "python-docx is required for report generation. Install `python-docx` to enable DOCX export."
        ) from _DOCX_IMPORT_ERROR

    document = Document()

    document.add_heading("Automated Essay Scoring Report", level=1)
    document.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    meta = document.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    rows = [
        ("Student", student_name or "-"),
        ("Submission ID", submission_id or "-"),
        ("Final Score", f"{evaluation.final_score}/6"),
        ("Similarity Risk", evaluation.similarity_risk),
    ]
    for i, (label, value) in enumerate(rows):
        meta.cell(i, 0).text = label
        meta.cell(i, 1).text = value

    document.add_heading("Prompt", level=2)
    document.add_paragraph(prompt or "-")

    document.add_heading("Essay", level=2)
    document.add_paragraph(essay_text or "-")

    document.add_heading("Rubric Breakdown", level=2)
    rub_table = document.add_table(rows=1, cols=2)
    rub_table.style = "Table Grid"
    hdr = rub_table.rows[0].cells
    hdr[0].text = "Dimension"
    hdr[1].text = "Score"
    for name, value in evaluation.rubric_breakdown.items():
        row = rub_table.add_row().cells
        row[0].text = name.capitalize()
        row[1].text = f"{value:.2f}"

    document.add_heading("Strengths", level=2)
    for item in evaluation.strengths:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("Weaknesses", level=2)
    for item in evaluation.weaknesses:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("Feedback", level=2)
    document.add_paragraph(evaluation.feedback)

    document.add_heading("Recommended Resources", level=2)
    for item in evaluation.recommendations:
        document.add_paragraph(item, style="List Bullet")

    if references:
        document.add_heading("Reference / Similarity Inputs", level=2)
        for ref in references:
            document.add_paragraph(ref[:600], style="List Bullet")

    out = BytesIO()
    document.save(out)
    out.seek(0)
    return out


def docx_filename(submission_id: str) -> str:
    slug = submission_id.strip() or "submission"
    return f"essay_report_{slug}.docx"


def build_pdf_report(
    submission_id: str,
    student_name: str,
    prompt: str,
    essay_text: str,
    evaluation: EssayEvaluation,
    references: Sequence[str] | None = None,
) -> BytesIO:
    if SimpleDocTemplate is None:
        raise RuntimeError(
            "reportlab is required for PDF export. Install `reportlab` to enable PDF export."
        ) from _PDF_IMPORT_ERROR

    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.6 * inch,
        leftMargin=0.6 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "AESTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#111827"),
        spaceAfter=10,
    )
    heading_style = ParagraphStyle(
        "AESHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=14,
        textColor=colors.HexColor("#111827"),
        spaceBefore=8,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "AESBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=12,
        textColor=colors.HexColor("#1f2937"),
    )

    story = []
    story.append(Paragraph("Intelligent Assignment Evaluator", title_style))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", body_style))
    story.append(Spacer(1, 0.14 * inch))

    meta = [
        ["Student", student_name or "-"],
        ["Submission ID", submission_id or "-"],
        ["Final Score", f"{evaluation.final_score}/6"],
        ["Similarity Risk", evaluation.similarity_risk],
    ]
    meta_table = Table(meta, colWidths=[1.4 * inch, 4.8 * inch])
    meta_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eff6ff")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#111827")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1d5db")),
                ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#e5e7eb")),
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 9.25),
                ("LEADING", (0, 0), (-1, -1), 12),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(meta_table)
    story.append(Spacer(1, 0.18 * inch))

    sections = [
        ("Prompt", prompt or "-"),
        ("Essay", essay_text or "-"),
        ("Feedback", evaluation.feedback),
    ]
    for heading, text in sections:
        story.append(Paragraph(heading, heading_style))
        story.append(Paragraph(text.replace("\n", "<br/>") if text else "-", body_style))
        story.append(Spacer(1, 0.08 * inch))

    story.append(Paragraph("Rubric Breakdown", heading_style))
    rub_rows = [["Dimension", "Score"]] + [
        [name.capitalize(), f"{value:.2f}"] for name, value in evaluation.rubric_breakdown.items()
    ]
    rub_table = Table(rub_rows, colWidths=[3.4 * inch, 1.4 * inch])
    rub_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eff6ff")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#111827")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1d5db")),
                ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#e5e7eb")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("LEADING", (0, 0), (-1, -1), 12),
            ]
        )
    )
    story.append(rub_table)
    story.append(Spacer(1, 0.12 * inch))

    story.append(Paragraph("Strengths", heading_style))
    story.append(Paragraph("<br/>".join(f"• {item}" for item in evaluation.strengths), body_style))
    story.append(Spacer(1, 0.08 * inch))

    story.append(Paragraph("Weaknesses", heading_style))
    story.append(Paragraph("<br/>".join(f"• {item}" for item in evaluation.weaknesses), body_style))
    story.append(Spacer(1, 0.08 * inch))

    story.append(Paragraph("Recommended Resources", heading_style))
    story.append(Paragraph("<br/>".join(f"• {item}" for item in evaluation.recommendations), body_style))

    if references:
        story.append(Spacer(1, 0.08 * inch))
        story.append(Paragraph("Reference / Similarity Inputs", heading_style))
        story.append(Paragraph("<br/>".join(f"• {ref[:280]}" for ref in references), body_style))

    document.build(story)
    buffer.seek(0)
    return buffer


def pdf_filename(submission_id: str) -> str:
    slug = submission_id.strip() or "submission"
    return f"essay_report_{slug}.pdf"
